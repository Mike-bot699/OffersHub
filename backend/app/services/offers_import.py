"""Importer for the historical "Offers Monitoring" workbook.

Supported inputs:
- .xlsx with sheets "Events", "Fix Codes", "Store Discounts" (others are reported and ignored)
- .csv export of the "Events" sheet

Every row is normalised and checked; anything doubtful becomes a visible `data_issue`
on the imported event rather than being silently corrected. Running with dry_run=True
returns the full preview without writing.
"""

import csv
import hashlib
import io
import logging
import re
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import UTC, date
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import ConflictError, ValidationError
from app.db.base import utcnow
from app.models import Coupon, DataNote, ImportBatch, Offer, PromoEvent, Store
from app.models.enums import DataSource, VerificationStatus
from app.offers.classify import classify_offer
from app.offers.status import set_status
from app.services.audit import audit
from app.services.parsing import (
    detect_currency,
    norm_key,
    parse_date,
    parse_number,
    parse_percent,
    store_code,
    strip_decoration,
)

log = logging.getLogger(__name__)

EVENT_COLUMNS = {
    "event name": "name", "store": "store", "year": "year", "start date": "start", "end date": "end",
    "duration": "duration", "offer": "offer", "type": "type", "coupon": "coupon", "threshold": "threshold",
    "offer activation": "activation", "bonus": "bonus", "displayed": "display", "offers orders": "offer_orders",
    "total orders": "total_orders", "%": "pct", "sales amount": "sales", "total sent": "emails_sent",
    "order per email": "orders_per_email", "bonus give away": "bonus_giveaway", "notes": "notes",
    # Marketing calendar (Events sheet) columns
    "event description": "description", "description": "description", "days": "duration",
    "collection": "collection", "level": "collection", "article": "articles", "banners message": "banner_message",
    "banner message": "banner_message", "landing page": "landing_page", "activation": "activation",
    "status": "calendar_status",
}
# The calendar's Status column mixes email states (Sent, Off) with offer states. Only the offer
# states are used; anything else leaves the status to other evidence.
CALENDAR_STATUS = {"live": "live", "paused": "paused", "cancelled": "cancelled", "canceled": "cancelled",
                   "scheduled": "scheduled", "completed": "completed", "ended": "completed"}
CODE_RE = re.compile(r"^[A-Z0-9][A-Z0-9_-]{2,}$")


def parse_codes(raw: str) -> list[str]:
    """Codes from a cell such as "SPOOKLAB (new) | OWLPASS (existing 15/20%)" or "N/A (automatic)".
    Notes in parentheses are dropped; only code-shaped tokens are kept."""
    text = re.sub(r"\([^)]*\)", " ", raw or "")
    out = []
    for tok in re.split(r"[|,/\n+;]", text):
        tok = tok.strip().upper()
        if CODE_RE.match(tok) and tok not in ("N/A", "NONE", "AUTOMATIC", "AUTO") and tok not in out:
            out.append(tok)
    return out

DISPLAY_MAP = {"full displayed": "full_site", "fully displayed": "full_site", "email only": "email_only",
               "content email": "email_only"}


def _header_map(headers: list[Any]) -> dict[int, str]:
    out: dict[int, str] = {}
    for i, h in enumerate(headers):
        key = strip_decoration(str(h or "")).lower().replace("-", " ").strip()
        if key == "%":
            out[i] = "pct"
            continue
        for prefix, name in EVENT_COLUMNS.items():
            if key.startswith(prefix) and name not in out.values():
                out[i] = name
                break
    return out


@dataclass
class ImportReport:
    filename: str
    sha256: str
    dry_run: bool
    sheets_found: list[str] = field(default_factory=list)
    sheets_ignored: list[str] = field(default_factory=list)
    events_parsed: int = 0
    events_created: int = 0
    events_updated: int = 0
    offers_created: int = 0
    coupons_upserted: int = 0
    rows_skipped: int = 0
    stores_existing: list[str] = field(default_factory=list)
    stores_created: list[str] = field(default_factory=list)
    store_settings_updated: list[str] = field(default_factory=list)
    issues: list[dict] = field(default_factory=list)  # {"sheet","row","store","event","issue"}
    issue_counts: dict[str, int] = field(default_factory=dict)
    batch_id: int | None = None

    def add_issue(self, sheet: str, row: int, issue: str, *, store: str | None = None, event: str | None = None):
        self.issues.append({"sheet": sheet, "row": row, "store": store, "event": event, "issue": issue})
        kind = re.sub(r"'[^']*'|\([^)]*\)|\d[\d.,/%-]*|[$€£]", "…", issue.split(" — ")[0])
        kind = re.sub(r"(…\s*)+", "… ", kind).strip()[:80]
        self.issue_counts[kind] = self.issue_counts.get(kind, 0) + 1


def _rows_from_file(content: bytes, filename: str) -> dict[str, list[list[Any]]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise ValidationError(f"Could not read workbook: {exc}") from exc
        sheets = {}
        for ws in wb.worksheets:
            sheets[ws.title] = [list(r) for r in ws.iter_rows(values_only=True)]
        return sheets
    if name.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        return {"Events": [row for row in csv.reader(io.StringIO(text))]}
    raise ValidationError("Unsupported file type. Upload the workbook as .xlsx or the Events sheet as .csv.")


def _find_sheet(sheets: dict[str, list], wanted: str) -> str | None:
    for title in sheets:
        if norm_key(title) == norm_key(wanted):
            return title
    return None


def import_offers_workbook(db: Session, content: bytes, filename: str, *, dry_run: bool = True,
                           create_missing_stores: bool = True, force: bool = False,
                           today: date | None = None) -> ImportReport:
    today = today or date.today()
    sha = hashlib.sha256(content).hexdigest()
    report = ImportReport(filename=filename, sha256=sha, dry_run=dry_run)
    if not dry_run and not force:
        prior = db.scalar(select(ImportBatch).where(ImportBatch.file_sha256 == sha,
                                                    ImportBatch.kind == "offers_workbook",
                                                    ImportBatch.status == "completed"))
        if prior:
            raise ConflictError(
                f"This exact file was already imported on {prior.created_at:%Y-%m-%d %H:%M}. "
                "Re-importing an edited version is fine; use force to re-import the identical file.",
                details={"batch_id": prior.id},
            )

    sheets = _rows_from_file(content, filename)
    report.sheets_found = list(sheets)
    events_sheet = _find_sheet(sheets, "Events")
    codes_sheet = _find_sheet(sheets, "Fix Codes")
    discounts_sheet = _find_sheet(sheets, "Store Discounts")
    used = {events_sheet, codes_sheet, discounts_sheet}
    report.sheets_ignored = [s for s in sheets if s not in used]
    if events_sheet is None:
        raise ValidationError("No 'Events' sheet found in the file.")

    batch = None
    if not dry_run:
        batch = ImportBatch(kind="offers_workbook", filename=filename, file_sha256=sha, status="running",
                            created_at=utcnow())
        db.add(batch)
        db.flush()

    stores = _resolve_stores(db, sheets, events_sheet, codes_sheet, report, create_missing_stores, dry_run)
    _import_events(db, sheets[events_sheet], events_sheet, stores, report, batch, dry_run, today)
    if codes_sheet:
        _import_codes(db, sheets[codes_sheet], codes_sheet, stores, report, dry_run)
    if discounts_sheet:
        _import_store_discounts(db, sheets[discounts_sheet], discounts_sheet, stores, report, dry_run)
    if not dry_run:
        from app.offers.status import sync_store_offers

        for st in stores.values():
            sync_store_offers(db, st, today=today)  # technical status from store evidence, if any

    if batch is not None:
        batch.status = "completed"
        batch.rows_total = report.events_parsed + report.rows_skipped
        batch.rows_imported = report.events_created + report.events_updated
        batch.rows_skipped = report.rows_skipped
        batch.warnings = report.issues[:2000]
        report.batch_id = batch.id
        audit(db, "import.offers_workbook", "import_batch", batch.id,
              details={"filename": filename, "events_created": report.events_created,
                       "events_updated": report.events_updated, "issues": len(report.issues)})
        db.commit()
    else:
        db.rollback()
    return report


def _resolve_stores(db, sheets, events_sheet, codes_sheet, report, create_missing, dry_run) -> dict[str, Store]:
    codes: Counter = Counter()
    currency_votes: dict[str, Counter] = defaultdict(Counter)
    for sheet in (events_sheet, codes_sheet):
        if not sheet:
            continue
        rows = sheets[sheet]
        if not rows:
            continue
        hdr = [strip_decoration(str(h or "")).lower() for h in rows[0]]
        if "store" not in hdr:
            continue
        si = hdr.index("store")
        for r in rows[1:]:
            if si < len(r) and r[si]:
                code = store_code(str(r[si]))
                if code:
                    codes[code] += 1
                    for cell in r:
                        cur = detect_currency(str(cell)) if isinstance(cell, str) else None
                        if cur:
                            currency_votes[code][cur] += 1
    existing = {s.code: s for s in db.scalars(select(Store))}
    out: dict[str, Store] = {}
    for code in sorted(codes):
        if code in existing:
            if existing[code].is_demo:
                report.add_issue("-", 0, f"store code {code} belongs to a demo store; rows skipped", store=code)
                continue
            out[code] = existing[code]
            report.stores_existing.append(code)
        elif create_missing:
            currency = (currency_votes[code].most_common(1) or [("USD", 0)])[0][0]
            st = Store(code=code, name=code, platform="shopify", currency=currency, status="active",
                       notes="Created from offers workbook import; connect the store in Integrations to verify.")
            if not dry_run:
                db.add(st)
                db.flush()
            out[code] = st
            report.stores_created.append(code)
        else:
            report.add_issue("-", 0, f"unknown store code {code}; rows skipped", store=code)
    return out


def parse_level(raw) -> tuple[str | None, list[str]]:
    """Where an offer applies, from the calendar's Collection column.
    Storewide / All -> store; a /products/ URL -> that product; anything else -> collection(s)."""
    text = strip_decoration(str(raw or "")).strip()
    if not text or text in ("-", "—"):
        return None, []
    low = text.lower()
    if low in ("all", "store", "whole store") or low.startswith(("storewide", "sitewide")):
        return "store", []
    parts = [p.strip() for p in re.split(r"[\n,]+", text) if p.strip()]
    targets, level = [], "collection"
    for part in parts:
        m = re.search(r"/(collections|products)/([^/?#\s]+)", part)
        if m:
            if m.group(1) == "products":
                level = "product"
            targets.append(m.group(2).replace("-", " ").title())
        else:
            targets.append(part)
    return level, targets


def _cell(row: list, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    v = row[idx]
    return v.strip() if isinstance(v, str) else v


def _import_events(db, rows, sheet, stores, report: ImportReport, batch, dry_run, today):
    if not rows:
        return
    hmap = _header_map(rows[0])
    col = {name: i for i, name in hmap.items()}
    missing = [c for c in ("name", "store", "start", "end", "offer") if c not in col]
    if missing:
        raise ValidationError(f"Events sheet is missing expected columns: {', '.join(missing)}")
    offers_cache: dict[str, Offer] = {}
    sales_seen: dict[float, tuple[int, str, str]] = {}

    for rn, row in enumerate(rows[1:], start=2):
        if not any(c not in (None, "") for c in row):
            continue
        g = lambda k: _cell(row, col.get(k))
        raw_name = g("name")
        code = store_code(str(g("store") or ""))
        if not raw_name or not code:
            report.rows_skipped += 1
            report.add_issue(sheet, rn, "row skipped: missing event name or store")
            continue
        store = stores.get(code)
        if store is None:
            report.rows_skipped += 1
            continue
        name = strip_decoration(str(raw_name))
        issues: list[str] = []
        year_val, _ = parse_number(g("year"), kind="count")
        year = int(year_val) if year_val else None
        start, i1 = parse_date(g("start"), default_year=year)
        end, i2 = parse_date(g("end"), default_year=year)
        issues += [i for i in (i1, i2) if i]
        if start and end and end < start:
            issues.append(f"end date {end} is before start date {start}")
        if start and year and start.year != year and not (abs(start.year - year) == 1 and start.month in (1, 12)):
            issues.append(f"start date year {start.year} differs from Year column {year}")

        offer_text = strip_decoration(str(g("offer") or "")) or None
        bonus = strip_decoration(str(g("bonus") or "")) or None
        coupon_raw = str(g("coupon") or "").strip()
        coupons = parse_codes(coupon_raw)

        thr_raw = str(g("threshold") or "").strip()
        thr_amount = None
        thr_text = thr_raw or None
        if thr_raw.lower() in ("no minimum", "n/a", "none", "0.0"):
            thr_raw, thr_text = "", "none"
        elif thr_raw and not any(ch.isdigit() for ch in thr_raw) and thr_raw not in ("-", "—"):
            thr_raw = ""  # a description such as "Number of items in cart": kept as text, not a number
        if thr_raw and thr_raw not in ("-", "—"):
            parts = [p for p in thr_raw.replace("/", " / ").split(" / ") if p.strip()]
            if len(parts) > 1:
                thr_amount, _ = parse_number(parts[0])
                if any(ch.isalpha() for ch in thr_raw):
                    issues.append(f"threshold '{thr_raw}' mixes quantity and amount; stored as text")
                    thr_amount = None
            else:
                thr_amount, ti = parse_number(thr_raw)
                if ti:
                    issues.append(f"threshold: {ti}")
            if thr_amount == 0:
                thr_amount = None
                thr_text = "none"

        offer_orders, _ = parse_number(g("offer_orders"), kind="count")
        total_orders, _ = parse_number(g("total_orders"), kind="count")
        sales, si = parse_number(g("sales"))
        if si:
            issues.append(f"sales amount: {si}")
        emails, _ = parse_number(g("emails_sent"), kind="count")
        pct, _ = parse_percent(g("pct"))
        if offer_orders is not None and total_orders:
            if offer_orders > total_orders:
                issues.append(f"offer orders ({offer_orders:.0f}) exceed total orders ({total_orders:.0f})")
            elif pct is not None and abs(offer_orders / total_orders - pct) > 0.005:
                issues.append(f"recorded share {pct:.2%} differs from offer/total orders "
                              f"{offer_orders / total_orders:.2%}")
        row_cur = detect_currency(str(g("sales") or "")) or detect_currency(thr_raw)
        if row_cur and store.currency and row_cur != store.currency:
            issues.append(f"amounts written in {row_cur} but store currency is {store.currency}")
        if sales and sales > 0:
            prior = sales_seen.get(round(sales, 2))
            if prior and (prior[1] != code or prior[2] != norm_key(name)):
                issues.append(f"sales amount {sales:,.2f} is identical to row {prior[0]} — possible copy/paste error")
            else:
                sales_seen[round(sales, 2)] = (rn, code, norm_key(name))
        if offer_orders and sales is not None and sales == 0:
            issues.append("offer orders recorded but sales amount is zero")

        notes = strip_decoration(str(g("notes") or "")) or None
        # Offer status is set only where the workbook gives evidence; dates alone never make an
        # offer Live or Completed. Everything else is left "not set" for the user to confirm.
        has_results = bool(offer_orders) or bool(sales)
        cal_status = CALENDAR_STATUS.get(norm_key(strip_decoration(str(g("calendar_status") or ""))))
        if notes and "not launched" in notes.lower():
            status, status_note = "cancelled", "Workbook note says it was not launched"
        elif cal_status:
            status, status_note = cal_status, f"Calendar status: {strip_decoration(str(g('calendar_status')))}"
        elif not start and not offer_text:
            status, status_note = None, None
            issues.append("no dates or offer recorded (placeholder row)")
        elif end and end < today and has_results:
            status, status_note = "completed", "Workbook records results for this offer"
        elif start and start > today:
            # Future dates: the offer is planned, so it enters the plan as Scheduled.
            status, status_note = "scheduled", "Future dates in the workbook"
        else:
            status, status_note = None, None
            if not start:
                issues.append("missing start/end date")

        activation_raw = str(g("activation") or "").lower()
        activation = "automatic" if activation_raw.startswith("auto") else "manual_code" if activation_raw else None
        display_raw = norm_key(g("display"))
        display = next((v for k, v in DISPLAY_MAP.items() if k in display_raw), None) if display_raw else None
        channels = ["email"] if display == "email_only" else (["site", "email"] if display == "full_site" else [])

        level, level_targets = parse_level(g("collection"))
        landing = str(g("landing_page") or "").strip()
        landing = landing if landing.lower().startswith(("http://", "https://", "www.")) else None
        cls = classify_offer(offer_text, bonus, coupon_raw)
        key_src = f"{code}|{norm_key(name)}|{start or year}|{'/'.join(coupons)}"
        nkey = hashlib.sha1(key_src.encode()).hexdigest()[:24]
        report.events_parsed += 1
        for issue in issues:
            report.add_issue(sheet, rn, issue, store=code, event=name)
        if dry_run:
            continue

        offer = None
        if offer_text:
            okey = hashlib.sha1(f"{norm_key(offer_text)}|{thr_amount}".encode()).hexdigest()[:24]
            offer = offers_cache.get(okey) or db.scalar(select(Offer).where(Offer.natural_key == okey))
            if offer is None:
                offer = Offer(name=offer_text[:300], offer_type=cls["offer_type"],
                              mechanic={**cls["mechanic"], **({"threshold_amount": thr_amount} if thr_amount else {})},
                              description=bonus, status="active", source=DataSource.SPREADSHEET_IMPORT,
                              natural_key=okey, eligible_scope={"type": "unknown"})
                db.add(offer)
                db.flush()
                report.offers_created += 1
            offers_cache[okey] = offer

        ev = db.scalar(select(PromoEvent).where(PromoEvent.store_id == store.id, PromoEvent.natural_key == nkey))
        if ev is None:
            ev = PromoEvent(store_id=store.id, natural_key=nkey, name=name)
            db.add(ev)
            report.events_created += 1
        else:
            report.events_updated += 1
        ev.name = name
        ev.offer_id = offer.id if offer else None
        ev.offer_text = offer_text
        ev.offer_type = cls["offer_type"] if offer_text else None
        ev.start_date, ev.end_date = start, end
        if ev.status_source in (None, "import"):  # never overwrite a status a person or the store has set
            if ev.id is None:
                ev.status = None
                db.flush()
            set_status(db, ev, status, source="import", actor="workbook import", note=status_note)
        ev.channels = channels
        ev.display = display
        ev.activation = activation
        ev.coupon_codes = coupons
        ev.threshold_amount = thr_amount
        ev.threshold_text = thr_text
        ev.bonus = bonus
        if col.get("description") is not None:
            ev.description = strip_decoration(str(g("description") or "")) or None
        if col.get("collection") is not None:
            ev.level, ev.level_targets = level, level_targets
        if col.get("landing_page") is not None:
            ev.landing_page = landing
        if col.get("articles") is not None:
            ev.articles = strip_decoration(str(g("articles") or "")) or None
        if col.get("banner_message") is not None:
            ev.banner_message = str(g("banner_message") or "").strip() or None
        ev.offer_orders = int(offer_orders) if offer_orders is not None else None
        ev.total_orders = int(total_orders) if total_orders is not None else None
        ev.offer_sales = sales
        ev.emails_sent = int(emails) if emails is not None else None
        ev.results_source = DataSource.SPREADSHEET_IMPORT
        ev.notes = notes
        ev.data_issues = issues
        ev.source = DataSource.SPREADSHEET_IMPORT
        ev.source_ref = f"{sheet}!row {rn}"
        ev.import_batch_id = batch.id if batch else None
        if ev.verification_status in (None, ""):
            ev.verification_status = VerificationStatus.DOCUMENTED_ONLY
    db.flush()


def _import_codes(db, rows, sheet, stores, report: ImportReport, dry_run):
    if not rows:
        return
    hdr = [norm_key(h) for h in rows[0]]

    def idx(*names):
        for n in names:
            for i, h in enumerate(hdr):
                if h.startswith(n):
                    return i
        return None

    ci, si = idx("coupon"), idx("store")
    if ci is None or si is None:
        report.add_issue(sheet, 1, "Fix Codes sheet lacks Coupon/Store columns; skipped")
        return
    di, mi, ti, sti = idx("discount"), idx("min pu"), idx("times"), idx("status")
    sdi, edi, oi, chi, evi = idx("start"), idx("end"), idx("offer"), idx("channel"), idx("event")
    now = utcnow()
    for rn, row in enumerate(rows[1:], start=2):
        code = str(_cell(row, ci) or "").strip().upper()
        scode = store_code(str(_cell(row, si) or ""))
        if not code or not scode or scode not in stores:
            continue
        store = stores[scode]
        disc = str(_cell(row, di) or "")
        if "%" in disc:
            vtype, val = "percent", parse_number(disc.replace("%", ""))[0]
        elif "off" in disc.lower():
            vtype, val = "fixed", parse_number(disc)[0]
        else:
            vtype, val = None, None
        min_raw = str(_cell(row, mi) or "")
        min_sub, min_qty = None, None
        if "item" in min_raw.lower():
            q = parse_number(min_raw.split("+")[0], kind="count")[0]
            min_qty = int(q) if q else None
        elif min_raw and min_raw.lower() != "none":
            min_sub = parse_number(min_raw)[0]
        used = parse_number(_cell(row, ti), kind="count")[0]
        status = str(_cell(row, sti) or "").lower() or None
        sd, _ = parse_date(_cell(row, sdi))
        ed, _ = parse_date(_cell(row, edi))
        report.coupons_upserted += 1
        if dry_run:
            continue
        c = db.scalar(select(Coupon).where(Coupon.store_id == store.id, Coupon.source == DataSource.SPREADSHEET_IMPORT,
                                           Coupon.code == code))
        if c is None:
            c = Coupon(store_id=store.id, source=DataSource.SPREADSHEET_IMPORT, code=code)
            db.add(c)
        from datetime import datetime, time

        c.title = strip_decoration(str(_cell(row, oi) or "")) or None
        c.value_type, c.value = vtype, val
        c.min_subtotal, c.min_quantity = min_sub, min_qty
        c.usage_count = int(used) if used is not None else None
        c.status = status
        c.starts_at = datetime.combine(sd, time(), tzinfo=UTC) if sd else None
        c.ends_at = datetime.combine(ed, time(23, 59), tzinfo=UTC) if ed else None
        c.channel = strip_decoration(str(_cell(row, chi) or "")) or None
        c.raw = {"event": strip_decoration(str(_cell(row, evi) or "")), "sheet_row": rn,
                 "note": "usage count as recorded in the spreadsheet, not verified"}
        c.last_synced_at = now
    db.flush()


def _import_store_discounts(db, rows, sheet, stores, report: ImportReport, dry_run):
    if len(rows) < 2:
        return
    hdr = [norm_key(h) for h in rows[0]]
    try:
        code_i = hdr.index("code")
    except ValueError:
        report.add_issue(sheet, 1, "Store Discounts sheet has no Code column; skipped")
        return
    for rn, row in enumerate(rows[1:], start=2):
        code = store_code(str(_cell(row, code_i) or ""))
        if not code or code not in stores:
            continue
        new_pct = parse_percent(_cell(row, code_i + 1))[0]
        new_thr = parse_number(_cell(row, code_i + 2))[0]
        ret_pct = parse_percent(_cell(row, code_i + 3))[0]
        ret_thr = parse_number(_cell(row, code_i + 4))[0]
        ship_thr = parse_number(_cell(row, code_i + 5))[0]
        report.store_settings_updated.append(code)
        if dry_run:
            continue
        store = stores[code]
        cs = dict(store.commercial_settings or {})
        cs["new_customer_discount_pct"] = new_pct
        cs["new_customer_discount_threshold"] = new_thr or None
        cs["returning_customer_discount_pct"] = ret_pct
        cs["returning_customer_discount_threshold"] = ret_thr or None
        cs["free_shipping_threshold"] = ship_thr or None
        store.commercial_settings = cs
        db.add(DataNote(store_id=store.id, entity_type="store", entity_id=store.id, created_at=utcnow(),
                        note=f"Standing discounts and free-shipping threshold imported from '{sheet}' row {rn} "
                             f"of the offers workbook (as documented, not verified against the store)."))
